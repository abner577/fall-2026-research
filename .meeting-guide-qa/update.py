from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
import zipfile
from lxml import etree

root=Path(__file__).resolve().parent
path=root.parent/'preparation'/'mentor-meeting-guide.docx'
doc=Document(path)
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
def body_text(document):
    return [''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in etree.fromstring(document._element.xml.encode()).xpath('//w:body//w:p',namespaces=ns)]
before=body_text(doc)
source=(root/'source.md').read_text(encoding='utf-8').splitlines()
new=[line for line in source if line.startswith(('Before walking through','Explain RAG in plain','Explain how that information'))]
anchor=next(p for p in doc.paragraphs if p.text.startswith('Start with:'))
for text in new:
    element=OxmlElement('w:p'); anchor._p.addnext(element)
    p=Paragraph(element,anchor._parent); p.style=doc.styles['Normal']; p.add_run(text); anchor=p
gather=next(p for p in doc.paragraphs if p.text.startswith('Gather context:'))
for r in list(gather.runs):r._r.getparent().remove(r._r)
gather.add_run('Gather context:').bold=True
gather.add_run(" The app collects the findings and surrounding code, then retrieves relevant original CWE passages from the knowledge base to include in the LLM's input.")
expected=before.copy()
idx=next(i for i,t in enumerate(expected) if t.startswith('Start with:'))
expected[idx+1:idx+1]=new
idx=next(i for i,t in enumerate(expected) if t.startswith('Gather context:'))
expected[idx]=gather.text
assert body_text(doc)==expected, 'Unexpected change outside requested edit'
path=root.parent/'preparation'/'mentor-meeting-guide-updated.docx'
doc.save(path)
reopened=Document(path)
assert body_text(reopened)==expected
with zipfile.ZipFile(path) as z:assert z.testzip() is None
print('Updated first section; verified all other wording is unchanged and DOCX package is valid.')
