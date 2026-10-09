from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

root=Path(__file__).resolve().parent
doc=Document()
s=doc.sections[0]
s.page_width=Inches(8.5); s.page_height=Inches(11)
s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1)
s.header_distance=s.footer_distance=Inches(.492)
# compact_reference_guide; no added title/header to preserve source content.
for name,size,color,before,after in [('Normal',11,'000000',0,6),('Heading 1',16,'2E74B5',18,10),('Heading 2',13,'2E74B5',14,7),('Heading 3',12,'1F4D78',10,5),('List Bullet',11,'000000',0,4)]:
    st=doc.styles[name]
    st.font.name='Calibri'; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color)
    pf=st.paragraph_format; pf.space_before=Pt(before); pf.space_after=Pt(after); pf.line_spacing=1.25
    pf.widow_control=True
    if name.startswith('Heading'): pf.keep_with_next=True

def numbering(abstract_id,num_id,fmt,text):
    a=OxmlElement('w:abstractNum'); a.set(qn('w:abstractNumId'),str(abstract_id))
    multi=OxmlElement('w:multiLevelType'); multi.set(qn('w:val'),'singleLevel'); a.append(multi)
    lvl=OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt',fmt),('lvlText',text),('lvlJc','left')]:
        e=OxmlElement('w:'+tag); e.set(qn('w:val'),val); lvl.append(e)
    pp=OxmlElement('w:pPr'); tabs=OxmlElement('w:tabs'); tab=OxmlElement('w:tab'); tab.set(qn('w:val'),'num'); tab.set(qn('w:pos'),'540'); tabs.append(tab); pp.append(tabs)
    ind=OxmlElement('w:ind'); ind.set(qn('w:left'),'540'); ind.set(qn('w:hanging'),'270'); pp.append(ind); lvl.append(pp); a.append(lvl)
    doc.part.numbering_part.element.append(a)
    n=OxmlElement('w:num'); n.set(qn('w:numId'),str(num_id)); aid=OxmlElement('w:abstractNumId'); aid.set(qn('w:val'),str(abstract_id)); n.append(aid); doc.part.numbering_part.element.append(n)
numbering(40,40,'bullet','•'); numbering(41,41,'decimal','%1.')
def num(p,n):
    pp=p._p.get_or_add_pPr(); np=OxmlElement('w:numPr')
    il=OxmlElement('w:ilvl'); il.set(qn('w:val'),'0'); ni=OxmlElement('w:numId'); ni.set(qn('w:val'),str(n)); np.append(il); np.append(ni); pp.append(np)
def runs(p,text):
    for i,part in enumerate(re.split(r'\*\*(.*?)\*\*',text)):
        r=p.add_run(part); r.bold=(i%2==1)

lines=(root/'source.md').read_text(encoding='utf-8').splitlines()
expected=[]; i=0
while i<len(lines):
    line=lines[i].strip(); i+=1
    if not line: continue
    if line.startswith('|'):
        rows=[line]
        while i<len(lines) and lines[i].strip().startswith('|'): rows.append(lines[i].strip()); i+=1
        data=[[c.strip() for c in row.strip('|').split('|')] for row in rows if not re.fullmatch(r'[|:\-\s]+',row)]
        t=doc.add_table(rows=0, cols=2); t.autofit=False
        widths=[3000,6360]
        pr=t._tbl.tblPr
        for tag,attrs in [('tblW',{'type':'dxa','w':'9360'}),('tblInd',{'type':'dxa','w':'120'})]:
            e=pr.find(qn('w:'+tag))
            if e is None: e=OxmlElement('w:'+tag); pr.append(e)
            for k,v in attrs.items(): e.set(qn('w:'+k),v)
        margins=OxmlElement('w:tblCellMar')
        for side,val in [('top',80),('bottom',80),('start',120),('end',120)]:
            e=OxmlElement('w:'+side); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); margins.append(e)
        pr.append(margins)
        borders=OxmlElement('w:tblBorders')
        for side in ['top','left','bottom','right','insideH','insideV']:
            e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:color'),'CBD5E1'); borders.append(e)
        pr.append(borders)
        for child in list(t._tbl.tblGrid): t._tbl.tblGrid.remove(child)
        for w in widths:
            e=OxmlElement('w:gridCol'); e.set(qn('w:w'),str(w)); t._tbl.tblGrid.append(e)
        for idx,rowdata in enumerate(data):
            row=t.add_row()
            cant=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(cant)
            if idx==0:
                rep=OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(rep)
            for j,text in enumerate(rowdata):
                c=row.cells[j]; c.width=Inches(widths[j]/1440); c._tc.get_or_add_tcPr().find(qn('w:tcW')).set(qn('w:w'),str(widths[j]))
                p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(4); runs(p,text); expected.append(text)
                if idx==0:
                    for r in p.runs:r.bold=True
                    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EEF5'); c._tc.get_or_add_tcPr().append(sh)
        continue
    if line.startswith('## '):
        text=line[3:]; p=doc.add_paragraph(style='Heading 1'); num(p,41)
    elif line.startswith('- '):
        text=line[2:]; p=doc.add_paragraph(style='List Bullet'); num(p,40)
    else:
        text=line; p=doc.add_paragraph()
        if re.fullmatch(r'\*\*.*\*\*',text):p.paragraph_format.keep_with_next=True
    runs(p,text); expected.append(text.replace('**',''))

footer=s.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run(); r.font.size=Pt(9); r.font.color.rgb=RGBColor.from_string('666666')
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
out=root.parent/'preparation'/'mentor-meeting-guide.docx'
doc.save(out)
# Verify body text, including table cells, against all source blocks in order.
from lxml import etree
import zipfile
with zipfile.ZipFile(out) as z:
    xml=etree.fromstring(z.read('word/document.xml'))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
actual=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in xml.xpath('//w:body//w:p',namespaces=ns)]
actual=[x for x in actual if x]
assert actual==expected, 'Source wording mismatch'
print(f'Created {out}; verified {len(expected)} text blocks unchanged.')
